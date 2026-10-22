from __future__ import annotations

import csv
import html
import json
import os
from pathlib import Path

from openpyxl import Workbook

from .analytics import report
from .domain import EVENT_TYPES
from .media import export_clip, montage


def safe_cell(v):
    if isinstance(v, (dict, list)):
        v = json.dumps(v, ensure_ascii=False)
    # Spreadsheet formula injection must not execute user-supplied names.
    if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
        return "'" + v
    return v


def display(v):
    if v is None:
        return "нет данных"
    return f"{v:.1f}%" if isinstance(v, float) else str(v)


def excel_cell(value):
    value=safe_cell(value)
    if isinstance(value,str) and len(value)>32767:
        return "Значение превышает лимит ячейки XLSX. Полное значение сохранено в report.json (тот же объект и поле); данные не усечены в JSON."
    return value


def court_svg(events, calibrations=(), label=""):
    parts = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 760 400" role="img" aria-label="Карта событий, не эффективность">',
             '<desc>'+html.escape(label)+'</desc>',
             '<defs><pattern id="unknown" width="9" height="9" patternUnits="userSpaceOnUse"><path d="M0 9L9 0" stroke="#758391"/></pattern></defs>',
             '<rect x="20" y="20" width="720" height="360" fill="url(#unknown)"/>']
    for c in calibrations:
        if not c.get("verified"):
            continue
        poly = c.get("visible_floor", [])
        if len(poly) >= 3:
            points = " ".join(f"{20 + x * 40},{20 + y * 40}" for x, y in poly)
            parts.append(f'<polygon points="{points}" fill="#cce7d7" opacity=".7"/>')
    parts += ['<rect x="20" y="20" width="720" height="360" fill="none" stroke="#40596d" stroke-width="3"/>',
              '<path d="M260 20V380M380 20V380M500 20V380" stroke="#40596d" fill="none"/>']
    for e in events:
        xy = e.get("floor_point")
        if xy:
            x, y = xy
            parts.append(f'<a href="#event-{e["id"]}"><circle cx="{20 + x*40}" cy="{20 + y*40}" r="5" fill="#c44b2c"><title>{html.escape(EVENT_TYPES[e["type"]])}, {e["start"]:.3f} с</title></circle></a>')
    parts.append('</svg>')
    return "".join(parts)


def event_link(eid):
    return f'<a href="#event-{html.escape(eid)}">эпизод {html.escape(eid[:8])}</a>'


def validated_court_profiles(project,filters=None):
    """Conservative static floor coverage, rebuilt from source-scoped evidence."""
    from .calibration import validate_profile,map_floor
    filters=filters or {};sources={s["id"]:s for s in project.all("source")}
    # A static mask cannot describe coverage through an uncalibrated camera session.
    uncertain={w["source_id"] for w in project.all("geometry_warning") if w.get("state") in ("unknown","shifted")}
    uncertain.update(s["source_id"] for s in project.all("camera_session"))
    profiles=project.all("calibration");superseded={p.get("parent_id") for p in profiles}
    out=[]
    for profile in profiles:
        source=sources.get(profile.get("source_id"))
        if not source or profile.get("id") in superseded or source["id"] in uncertain:continue
        if filters.get("source_id") and filters["source_id"]!=source["id"]:continue
        try:
            checked=validate_profile(profile,source_id=source["id"],source_sha256=source["sha256"],
                image_size=[source.get("display_width",source["width"]),source.get("display_height",source["height"])])
            if not checked["verified"] or not checked.get("visible_image"):continue
            t=checked.get("reference_time",checked["start"])
            floor=[map_floor(checked,p,source["id"],t,on_floor=True)["floor_point"] for p in checked["visible_image"]]
            if all(p is not None for p in floor):out.append(checked|{"visible_floor":floor})
        except (ValueError,KeyError):continue
    return out


def observation_grid_svg(source,seconds=False,label=""):
    """Image-coordinate counts/time estimates; no pretend court geometry."""
    values=source["presence_seconds_grid"] if seconds else source["detection_grid"]
    maximum=max((v for row in values for v in row),default=0)
    rows,columns=source["grid_rows"],source["grid_columns"]
    width,height=720,720*source["height"]/source["width"]
    unit="оценка человеко-секунд" if seconds else "наблюдения детектора"
    parts=[f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width+40} {height+60}" role="img" aria-label="{unit}">',
           '<desc>'+html.escape(label)+'. '+unit+'; координаты изображения, не площадки</desc>',
           '<defs><pattern id="missing" width="8" height="8" patternUnits="userSpaceOnUse"><path d="M0 8L8 0" stroke="#a0abb5"/></pattern></defs>']
    for y,row in enumerate(values):
        for x,value in enumerate(row):
            left,top=20+x*width/columns,25+y*height/rows
            color=f"rgba(32,130,125,{0.15+0.8*value/maximum:.3f})" if value>0 and maximum else "url(#missing)"
            evidence=source.get("evidence_grid",[])
            pointer=evidence[y][x] if evidence else None
            t=pointer["representative_times"][0] if pointer and pointer.get("representative_times") else None
            if t is not None:parts.append(f'<a href="viewer.html#{html.escape(source["source_id"],quote=True)}:{t:.6f}">')
            caption=f"{value:.2f}" if seconds else str(value)
            parts.append(f'<rect x="{left}" y="{top}" width="{width/columns}" height="{height/rows}" fill="{color}" stroke="#f8f9fa" stroke-width=".6"><title>{caption} {unit}; ячейка ({x},{y})</title></rect>')
            if value>0:parts.append(f'<text x="{left+width/columns/2}" y="{top+height/rows/2+4}" text-anchor="middle" font-size="11" fill="#172f38">{caption}</text>')
            if t is not None:parts.append('</a>')
    parts.append(f'<text x="20" y="{height+48}" font-size="12">{source["width"]} × {source["height"]} пикселей · {unit}</text></svg>')
    return ''.join(parts)


def export_all(project, target=None, mode="automatic", filters=None, clips=False, before=2, after=3, cancel=None):
    target = Path(target or project.root / "exports" / mode)
    target.mkdir(parents=True, exist_ok=True)
    data = report(project, mode, filters)
    (target / "report.json").write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    (target / "project-events.json").write_text(json.dumps(project.snapshot(), ensure_ascii=False, indent=2), encoding="utf-8")
    metadata = {k:data[k] for k in ("mode","mode_label","dictionary_version","filters","completeness","limitations")}
    metadata["synthetic"] = data["match"]["synthetic"]
    (target / "events.metadata.json").write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding="utf-8")
    if clips:
        clip_files = []
        for e in data["events"]:
            clip_file = target / "clips" / f"{e['id']}.mp4"
            export_clip(project, e, clip_file, before, after, cancel, output_size=(1280,720))
            clip_files.append(clip_file)
        if clip_files:
            montage(clip_files,target/"montage.mp4",cancel)
            (target/"montage.json").write_text(json.dumps(metadata|{"event_ids":[e["id"] for e in data["events"]],
                       "before_s":before,"after_s":after,"output_size":[1280,720],"audio":False},ensure_ascii=False,indent=2),encoding="utf-8")
    fields = ["id", "source_id", "start", "end", "type", "team_id", "player_id", "participants", "source_zone",
              "target_zone", "direction", "outcome", "quality", "review", "evidence", "unknown_reasons", "model_version"]
    with (target / "events.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["report_mode","dictionary_version","filters","synthetic",*fields])
        for e in data["events"]:
            writer.writerow([mode,data["dictionary_version"],safe_cell(filters or {}),data["match"]["synthetic"],*[safe_cell(e.get(k)) for k in fields]])
    wb = Workbook()
    wb.remove(wb.active)
    for title, values in (("Игроки", data["players"]), ("Команды", data["teams"]), ("События", data["events"]),
                          ("Неопределённости", data["uncertainties"]), ("Очки", data["point_ledger"]),
                          ("Геометрия",data["geometry_warnings"]),("Сессии камеры",data["camera_sessions"]),
                          ("Временная модель",data["temporal_diagnostics"]),
                          ('Ручной счёт',[s for p in data['manual_score']['sets'] for s in p['snapshots']]),
                          ('Сверка счёта',[r|{'set_id':p['set_id']} for p in data['manual_score']['sets'] for r in p['reconciliation']]),
                          ('Конфликты счёта',data['manual_score']['conflicts']),
                          ('Подсказки личности',data['identity_hypotheses']),
                          ("Время наблюдений",data["observation_maps"]["players"]),
                          ("Наблюдения источников",data["observation_maps"]["sources"])):
        ws = wb.create_sheet(title)
        keys = list(dict.fromkeys(k for row in values for k in row))
        ws.append(keys or ["Нет данных"])
        for row in values:
            ws.append([excel_cell(row.get(k)) for k in keys])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    ws = wb.create_sheet("Правила и фильтры")
    for k, v in (("Режим", data["mode_label"]), ("Фильтры", filters or {}), ("Данные синтетические", data["match"]["synthetic"]),
                 ("Полнота", data["completeness"]), ("Словарь", "local-1, не FL и не официальный стандарт"),
                 ("K/TA", "Результативные / оценённые атаки; неизвестные исходы отдельно"),
                 ("(K-E)/TA", "E включает blocked ровно один раз; нулевой знаменатель = нет данных"),
                 ("Очки", "Один известный победитель розыгрыша = одно командное очко; участники блока не умножают очки"),
                 ("Side-out", "Выигранные при приёме / известные исходы при известном подающем"),
                 ("Обработка", data["processing_diagnostics"]), ("Пропуски", data["limitations"])):
        ws.append([k, excel_cell(v) if v is not None else "не определена"])
    wb.save(target / "statistics.xlsx")
    svg = court_svg(data["events"], validated_court_profiles(project,filters),label=data["mode_label"]+ ("; СИНТЕТИЧЕСКИЕ ДАННЫЕ" if data["match"]["synthetic"] else ""))
    (target / "court.svg").write_text(svg, encoding="utf-8")
    observation_sections=[]
    for index,source in enumerate(data["observation_maps"]["sources"]):
        for seconds,name in ((False,"detections"),(True,"presence")):
            map_svg=observation_grid_svg(source,seconds,label=data["mode_label"])
            (target/f"{name}-{index+1}.svg").write_text(map_svg,encoding="utf-8")
            observation_sections.append(f'<h3>{"Оценка времени" if seconds else "Число детекций"} · {html.escape(source["source_id"][:8])}</h3>'+map_svg)
        duration=source["estimated_person_seconds"]
        observation_sections.append(f'<p>Наблюдений: {source["detection_count"]}; кадров выборки: {source["people_sampled_frames"]}; '
            f'оценка человеко-секунд: {"нет данных" if duration is None else f"{duration:.3f}"}. '
            f'Полнота относительно матча не определена. {html.escape(source.get("scope",""))}</p>')
        observation_sections.append('<ul>'+''.join('<li>'+html.escape(x)+'</li>' for x in source["limitations"])+'</ul>')
    if not observation_sections:
        observation_sections=['<p>Нет данных для карт наблюдений. '+html.escape(' '.join(data["observation_maps"].get("limitations",[])))+'</p>']
    title = html.escape(data["match"]["title"])
    tables = []
    for label, rows in (("Игроки", data["players"]), ("Команды", data["teams"])):
        body = []
        for r in rows:
            m = r["attack_efficiency"]
            links = " ".join(event_link(eid) for eid in r["event_ids"])
            body.append(f'<tr><td>{html.escape(r["name"])}</td><td>{r["attacks_recorded"]}</td><td>{r["attack_kills"]}</td>'
                        f'<td>{r["attack_errors_including_blocked"]}</td><td>{display(m["value"])} ({m["numerator"]}/{m["denominator"]})</td>'
                        f'<td>{m["missing"]}</td><td>{links or "нет данных"}</td></tr>')
        tables.append(f'<h2>{label}</h2><table><thead><tr><th>Имя</th><th>Атаки записаны</th><th>K</th><th>E</th>'
                      '<th>Эффективность по оценённым</th><th>Исход неизвестен</th><th>Эпизоды</th></tr></thead><tbody>' + "".join(body) + '</tbody></table>')
    event_rows = []
    all_events = {e["id"]: e for e in project.events()}
    episode_ids = set(all_events)  # include uncertain candidates for inspection without requiring inspection
    for eid in sorted(episode_ids, key=lambda x: (all_events[x]["source_id"], all_events[x]["start"])):
        e = all_events[eid]
        clip = target / "clips" / f"{eid}.mp4"
        href = f"clips/{eid}.mp4" if clip.exists() else f"viewer.html#{e['source_id']}:{e['start']:.6f}"
        event_rows.append(f'<tr id="event-{eid}"><td><a href="{href}">{eid[:8]} · {e["start"]:.3f} с</a></td>'
                          f'<td>{html.escape(EVENT_TYPES[e["type"]])}</td><td>{html.escape(e.get("review", ""))}</td>'
                          f'<td>{html.escape(str(e.get("unknown_reasons", {})))}</td></tr>')
    conclusions = "".join(f'<li>{html.escape(c["text"])}. {"Малая выборка. " if c["small_sample"] else ""}' +
                          " ".join(event_link(eid) for eid in c["metric"]["episode_ids"]) + '</li>' for c in data["conclusions"])
    banner = "СИНТЕТИЧЕСКИЕ ДАННЫЕ — не оценка распознавания" if data["match"]["synthetic"] else data["mode_label"]
    from .roster import IDENTITY_MODES
    identity_note=IDENTITY_MODES[data['identity_mode']]
    identity_note+='; новые связи внешности остаются гипотезами, статистика игроков использует явные привязки.'
    score_rows=[];team_names={t['id']:t['name'] for t in data['teams']}
    for party in data['manual_score']['sets']:
        for s in party['snapshots']:
            label=' · '.join(f"{team_names.get(tid,tid)}: {value}" for tid,value in s['scores'].items())
            pointer=(f'<a href="viewer.html#{html.escape(s["source_id"],quote=True)}:{s["time"]:.6f}">{s["time"]:.3f} с</a>'
                if s.get('source_id') and s.get('time') is not None else 'без привязки ко времени')
            score_rows.append(f'<li>Партия {party["number"]}: {html.escape(label)} · {pointer} · '+('итог' if s['final'] else 'начало' if s['initial'] else 'отметка')+' · ручной ввод</li>')
        for interval in party['reconciliation']:
            if interval['unattributed_team_points'] and any(interval['unattributed_team_points'].values()):
                values=' · '.join(f"{team_names.get(tid,tid)}: {count}" for tid,count in interval['unattributed_team_points'].items())
                score_rows.append('<li>Очки между отметками без соответствующего известного розыгрыша: '+html.escape(values)+'. Авторы и причины не определены.</li>')
    score_rows.extend('<li>Противоречие: '+html.escape(c['reason'])+'</li>' for c in data['manual_score']['conflicts'])
    geometry_rows=[]
    for item in data["geometry_warnings"]+data["camera_sessions"]:
        t=item.get("time",item.get("start"));sid=item.get("source_id")
        pointer=(f'<a href="viewer.html#{html.escape(sid,quote=True)}:{t:.6f}">{t:.3f} с</a>'
                 if sid and t is not None else "время не определено")
        geometry_rows.append('<li>'+html.escape(str(sid or "источник не определён"))+' · '+pointer+' · '+
            html.escape(str(item.get("reason",item.get("state","геометрия недоступна"))))+'</li>')
    document = f'''<!doctype html><html lang="ru"><meta charset="utf-8"><title>{title}</title>
<style>body{{font:16px system-ui;margin:40px auto;max-width:1200px;background:#f3f5f7;color:#243444}}h1{{font-size:32px}}table{{border-collapse:collapse;width:100%;background:white}}td,th{{text-align:left;border-bottom:1px solid #ccd4da;padding:10px;vertical-align:top}}a{{color:#17607d}}.banner{{padding:18px;background:#ffe6b0}}svg{{max-width:760px}}pre{{white-space:pre-wrap}}li{{margin:8px 0}}</style>
<h1>{title}</h1><p class="banner">{html.escape(banner)}</p><p>{html.escape(data['mode_label'])}</p>
<p>Идентификация: {html.escape(identity_note)}</p>
<p>Фильтры: {html.escape(str(data['filters']))}. Полнота: {display(data['completeness'])}.
Исключено из основной таблицы: {data['excluded_events']}; неизвестная команда: {data['unknown_team_events']}.</p>
<h2>Ручной счёт</h2><ul>{''.join(score_rows) or '<li>Счёт не введён; обработка доступных наблюдений продолжается.</li>'}</ul>
<p>{html.escape(' '.join(data['manual_score']['limitations']))}</p>
{''.join(tables)}<h2>Описательные наблюдения</h2><ul>{conclusions or '<li>Нет данных для выводов.</li>'}</ul>
<h2>Карта событий</h2><p>Точки — подтверждённые места на полу. Штриховка — наблюдаемость не подтверждена.
Зелёный фон — пол в проверенном профиле, а не доказательство непрерывного наблюдения.
Плотность событий не является эффективностью или временем присутствия. Зоны без данных не получают нулевую успешность.</p>{svg}
<h2>Карты наблюдений и времени присутствия</h2><p>Это отдельные диагностические карты детекций в координатах видео.
Штриховка — нет данных в ячейке; отсутствие игрока не доказано. Оценка времени не заполняет длительные потери.</p>{''.join(observation_sections)}
<h2>Все эпизоды и неопределённости</h2><table><tr><th>Видео</th><th>Действие</th><th>Проверка</th><th>Причины пропусков</th></tr>{''.join(event_rows)}</table>
<h2>Правила и ограничения</h2><p>Словарь local-1. K/TA и (K−E)/TA рассчитаны только по оценённым исходам.
Blocked входит в E один раз. Командные очки берутся из известных результатов розыгрышей.
Блоковые участия и индивидуально зачтённые очки различаются. Полный набор полей, числители,
знаменатели, пропуски и ссылки сохранены в report.json и XLSX.</p><ul>{''.join('<li>'+html.escape(x)+'</li>' for x in data['limitations'])}</ul>
<h2>Диагностика обработки</h2><p>ball_missing — кадры выборки без обнаруженного мяча. Это не полнота детекции относительно эталона.</p>
<h3>Геометрия и смены ракурса</h3><ul>{''.join(geometry_rows) or '<li>Нет зарегистрированных предупреждений; точность этим не подтверждается.</li>'}</ul>
<pre>{html.escape(json.dumps(data['processing_diagnostics'],ensure_ascii=False,indent=2))}</pre>
<h3>Временной модуль действий</h3><pre>{html.escape(json.dumps(data['temporal_diagnostics'],ensure_ascii=False,indent=2))}</pre></html>'''
    (target / "report.html").write_text(document, encoding="utf-8")
    sources = {}
    for s in project.all("source"):
        try:
            video = project.resolve(s, proxy=True)
            try:
                video.relative_to(project.root)
                sources[s["id"]] = Path(os.path.relpath(video,target)).as_posix()
            except ValueError:
                sources[s["id"]] = video.as_uri()
        except FileNotFoundError:
            sources[s["id"]] = None
    script_sources = json.dumps(sources, ensure_ascii=False).replace("<", "\\u003c")
    (target / "viewer.html").write_text('''<!doctype html><html lang="ru"><meta charset="utf-8"><title>Эпизод</title>
<p>Локальное видео; время относительно первого видеокадра. Для HEVC/MOV используйте плеер приложения или создайте proxy.</p>
<video controls style="width:95%;max-height:85vh"></video><p id="status"></p><script>
const sources=''' + script_sources + '''; const [id,t]=location.hash.slice(1).split(':');const v=document.querySelector('video');
if(sources[id]){v.src=sources[id];v.addEventListener('loadedmetadata',()=>{v.currentTime=Number(t)||0;});}
else document.querySelector('#status').textContent='Исходник недоступен: восстановите путь в приложении и повторите экспорт';
</script></html>''', encoding="utf-8")
    return target
