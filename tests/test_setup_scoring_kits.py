import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import numpy as np
import pytest
from PySide6 import QtWidgets as W

from volleyvision.analytics import report
from volleyvision.exporting import export_all
from volleyvision.identity import appearance,second_pass
from volleyvision.kits import save_sample,set_our_team,hypotheses
from volleyvision.scoring import save_score,score_report,remove_score
from volleyvision.setup_ui import KitDialog,ScoreDialog
from volleyvision.store import Project


@pytest.fixture
def project(tmp_path):
    with Project(tmp_path/'project',create=True,synthetic=True) as p:
        for tid in ('A','B'):p.put('team',{'id':tid,'name':tid})
        p.put('player',{'id':'lib','team_id':'A','number':7,'role':'libero'})
        p.put('set',{'id':'s1','number':1})
        p.put('source',{'id':'v','path':'missing.mp4','sha256':'sha','duration':100,'offset':0,'width':100,'height':100,'display_width':100,'display_height':100})
        yield p


def rally(project,tid,time,rid=None):
    return project.put('rally',{'id':rid or f'r-{time}','set_id':'s1','source_id':'v','start':time-.2,'end':time,
        'winner_team_id':tid,'evidence':{'winner_team_id':'confirmed'}})


def test_manual_score_reconciles_gaps_without_player_or_event_credit(project,tmp_path):
    save_score(project,'s1',{'A':0,'B':0},'v',0,initial=True)
    rally(project,'A',1);rally(project,'B',2)
    save_score(project,'s1',{'A':2,'B':1},'v',3)
    data=report(project)
    interval=data['manual_score']['sets'][0]['reconciliation'][-1]
    assert interval['score_delta']=={'A':2,'B':1} and interval['unattributed_team_points']=={'A':1,'B':0}
    assert not data['manual_score']['conflicts'] and data['events']==[]
    assert all(p['credited_points']==0 for p in data['players'])
    out=export_all(project,tmp_path/'report')
    assert 'viewer.html#v:3.000000' in (out/'report.html').read_text(encoding="utf-8")
    from openpyxl import load_workbook
    assert {'Ручной счёт','Сверка счёта','Конфликты счёта'}<=set(load_workbook(out/'statistics.xlsx').sheetnames)
    assert report(project,filters={'player_id':'lib'})['manual_score']==data['manual_score']


def test_manual_score_edit_remove_history_conflict_and_final_without_time(project):
    save_score(project,'s1',{'A':0,'B':0},'v',0,initial=True)
    rally(project,'A',1);rally(project,'A',2)
    wrong=save_score(project,'s1',{'A':1,'B':0},'v',3)
    assert score_report(project)['conflicts']
    fixed=save_score(project,'s1',{'A':2,'B':0},'v',3,object_id=wrong['id'])
    assert fixed['revision']==2 and not score_report(project)['conflicts']
    final=save_score(project,'s1',{'A':25,'B':19},final=True)
    result=score_report(project)['sets'][0]
    assert result['winner_team_id']=='A' and result['final_scores']=={'A':25,'B':19}
    with pytest.raises(ValueError,match='уже задан'):save_score(project,'s1',{'A':25,'B':20},final=True)
    remove_score(project,final['id'])
    assert score_report(project)['sets'][0]['final_scores'] is None
    assert len(project.all('score_history'))==5
    for scores in ({'A':1},{'A':-1,'B':0},{'A':True,'B':0}):
        with pytest.raises(ValueError):save_score(project,'s1',scores,'v',2)
    with pytest.raises(ValueError):save_score(project,'s1',{'A':0,'B':0},'v',101)


def test_libero_form_stays_in_team_similar_opponent_is_unknown(project):
    red=np.full((100,100,3),[210,20,20],np.uint8);blue=np.full((100,100,3),[20,20,210],np.uint8)
    save_sample(project,'v',0,red,[10,10,80,80],'A')
    lib=save_sample(project,'v',0,blue,[10,10,80,80],'A','libero','lib')
    set_our_team(project,'A')
    track=project.put('track',{'id':'t','run_id':'run','source_id':'v','start':0,'end':1,'samples':[],'appearance':appearance(blue)})
    proposal=hypotheses(project,track)
    assert proposal['status']=='hypothesis' and proposal['proposed_team_id']=='A'
    assert proposal['candidates'][0]['kit_kind']=='libero'
    project.put('source',{'id':'part2','path':'missing2.mp4','sha256':'sha2','duration':100,'offset':100})
    assert hypotheses(project,track|{'source_id':'part2'})['proposed_team_id']=='A'
    # Same kit on an opponent prevents a team decision; neither becomes a player binding.
    save_sample(project,'v',0,blue,[10,10,80,80],'B')
    assert hypotheses(project,track)['proposed_team_id'] is None
    second_pass(project,'run')
    assert project.all('track_binding')==[] and project.all('kit_hypothesis')[0]['status']=='unknown'
    project.put('side',{'id':'side','team_id':'A','side':'right','source_id':'v','start':0,'end':100})
    assert project.get('match','match')['our_team_id']=='A'
    assert lib['player_id']=='lib' and lib['team_id']=='A'
    with pytest.raises(ValueError):save_sample(project,'v',0,red,[10,10,80,80],'B','libero','lib')


def test_overlapping_rallies_do_not_pass_score_reconciliation(project):
    save_score(project,'s1',{'A':0,'B':0},'v',0,initial=True)
    rally(project,'A',1,'r1');rally(project,'A',1,'r2')
    save_score(project,'s1',{'A':2,'B':0},'v',2)
    data=score_report(project)
    assert data['conflicts'] and not data['sets'][0]['known_rally_points_independent']
    assert data['sets'][0]['reconciliation'][-1]['unattributed_team_points'] is None


def test_offscreen_setup_samples_and_score_dialog_roundtrip(project):
    app=W.QApplication.instance() or W.QApplication([])
    rgb=np.full((100,100,3),[210,20,20],np.uint8)
    kit=KitDialog(project,project.get('source','v'),rgb)
    kit.ours.setCurrentIndex(kit.ours.findData('A'));kit.team.setCurrentIndex(kit.team.findData('A'))
    kit.kind.setCurrentIndex(kit.kind.findData('libero'));kit.player.setCurrentIndex(kit.player.findData('lib'))
    kit.click(10,10);kit.click(80,80);kit.save();assert len(project.all('kit_sample'))==1
    score=ScoreDialog(project,'v',3);score.boxes['A'].setValue(10);score.boxes['B'].setValue(8);score.save()
    assert len(project.all('score_snapshot'))==1
    score.table.selectRow(0);score.boxes['A'].setValue(11);score.save()
    assert len(project.all('score_snapshot'))==1 and project.all('score_snapshot')[0]['scores']['A']==11
    score.close();kit.close();app.processEvents()
