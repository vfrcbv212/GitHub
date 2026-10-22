import json
import os

os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import numpy as np
import pytest
from PySide6 import QtWidgets as W

from volleyvision.analytics import report
from volleyvision.calibration import map_floor
from volleyvision.exporting import export_all,excel_cell,validated_court_profiles
from volleyvision.gui import CalibrationDialog
from volleyvision.store import Project,dumps


@pytest.fixture
def project(tmp_path):
    with Project(tmp_path/'project',create=True,synthetic=True) as p:
        p.put('source',{'id':'source','path':'missing.mp4','sha256':'sample-hash','duration':2,
                       'width':640,'height':360,'display_width':640,'display_height':360,'offset':0})
        p.put('team',{'id':'A','name':'A'})
        for pid in ('A1','A2'):p.put('player',{'id':pid,'team_id':'A','number':int(pid[-1]),'name':None})
        yield p


def test_approximate_zone_gui_saves_reopens_without_invented_metres(project):
    app=W.QApplication.instance() or W.QApplication([])
    source=project.get('source','source');rgb=np.zeros((360,640,3),np.uint8)
    dialog=CalibrationDialog(project,source,rgb)
    dialog.mode.setCurrentIndex(dialog.mode.findData('zone'))
    dialog.zone_team.setCurrentIndex(dialog.zone_team.findData('A'));dialog.zone_number.setValue(4)
    for x,y in ((50,100),(250,100),(250,300),(50,300)):dialog.click(x,y)
    dialog.finish_polygon();assert len(dialog.zones)==1 and not dialog.current
    dialog.save();assert W.QDialog.result(dialog)==W.QDialog.DialogCode.Accepted
    saved=project.all('calibration')[-1]
    mapped=map_floor(saved,[100,150],'source',1,team_id='A',on_floor=True)
    assert mapped['zone']==4 and mapped['floor_point'] is None and saved['geometry']=='approximate'
    reopened=CalibrationDialog(project,source,rgb)
    assert reopened.zones==saved['zones']
    reopened.close();dialog.close();app.processEvents()


def test_saved_refinement_is_not_replaced_by_original_click_fit(project):
    app=W.QApplication.instance() or W.QApplication([])
    dialog=CalibrationDialog(project,project.get('source','source'),np.zeros((360,640,3),np.uint8))
    dialog.points=[[50,100],[590,100],[590,330],[50,330]];dialog.world=[[0,0],[18,0],[18,9],[0,9]]
    dialog.controls=[([320,215],[9,4.5])];dialog.check();assert dialog.result is not None
    candidate=np.array(dialog.result['matrix']);candidate[0,2]+=.02
    dialog.result['matrix']=candidate.tolist()
    dialog.save();assert W.QDialog.result(dialog)==W.QDialog.DialogCode.Accepted
    assert np.asarray(project.all('calibration')[-1]['matrix'])==pytest.approx(candidate)
    reopened=CalibrationDialog(project,project.get('source','source'),np.zeros((360,640,3),np.uint8))
    reopened.save();assert W.QDialog.result(reopened)==W.QDialog.DialogCode.Accepted
    assert np.asarray(project.all('calibration')[-1]['matrix'])==pytest.approx(candidate)
    reopened.close();dialog.close();app.processEvents()


def test_profile_scope_selected_at_displayed_time_is_preserved(project):
    app=W.QApplication.instance() or W.QApplication([])
    for pid,start,end in (('early',0,.8),('late',1,2)):
        project.put('calibration',{'id':pid,'source_id':'source','source_sha256':'sample-hash',
                    'geometry':'video','start':start,'end':end,'image_size':[640,360]})
    rgb=np.zeros((360,640,3),np.uint8)
    early=CalibrationDialog(project,project.get('source','source'),rgb,timestamp=.5)
    assert early.previous['id']=='early'
    early.save();saved=project.all('calibration')[-1]
    assert saved['parent_id']=='early' and saved['start']==0 and saved['end']==.8
    gap=CalibrationDialog(project,project.get('source','source'),rgb,timestamp=.9)
    assert not gap.previous and gap.profile_start.value()==.9 and gap.profile_end.value()==1
    gap.close();early.close();app.processEvents()


def test_legacy_draft_matrix_reopens_without_accuracy_fields(project):
    app=W.QApplication.instance() or W.QApplication([])
    project.put('calibration',{'id':'legacy','source_id':'source','source_sha256':'sample-hash',
        'verified':True,'matrix':[[.02,0,0],[0,.02,0],[0,0,1]],'start':0,'end':2})
    dialog=CalibrationDialog(project,project.get('source','source'),np.zeros((360,640,3),np.uint8))
    dialog.save();assert W.QDialog.result(dialog)==W.QDialog.DialogCode.Accepted
    saved=project.all('calibration')[-1]
    assert not saved['verified'] and saved['geometry']=='video'
    dialog.close();app.processEvents()


def test_presence_export_uses_own_units_filters_and_binding_cache(project,tmp_path):
    project.put('job',{'id':'run','source_id':'source','params':{'params':{'people_max_gap':.4}},'state':'complete'})
    project.put('track',{'id':'fragment','run_id':'run','source_id':'source','start':0,'end':.4})
    project.bind_track('run','fragment',0,.4,'A1')
    for t in (0,.2,.4):
        body={'time':t,'source_pts':int(t*100),'people_sampled':True,'balls':[],
              'people':[{'track_id':'fragment','bbox':[50,100,100,200],'class_id':1,'observed':True}]}
        project.db.execute('INSERT INTO observations VALUES(?,?,?,?)',('run','source',t,dumps(body)))
    project.db.commit()
    data=report(project);source=data['observation_maps']['sources'][0]
    assert source['detection_count']==3 and source['estimated_person_seconds']==pytest.approx(.4)
    assert next(p for p in data['players'] if p['id']=='A1')['attack_efficiency']['value'] is None
    out=export_all(project,tmp_path/'report')
    assert (out/'detections-1.svg').exists() and (out/'presence-1.svg').exists()
    assert 'viewer.html#source:0.000000' in (out/'presence-1.svg').read_text(encoding="utf-8")
    assert (out/'presence-1.svg').read_text(encoding="utf-8").count('<a href=')==1
    assert report(project,'confirmed')['observation_maps']['sources']==[]
    assert report(project,filters={'type':'attack'})['observation_maps']['sources']==[]
    assert report(project,filters={'rally_id':'nonexistent'})['observation_maps']['sources']==[]
    project.bind_track('run','fragment',0,.4,'A2')
    updated=report(project)['observation_maps']['players']
    assert {p['player_id'] for p in updated}=={'A2'}
    assert report(project,filters={'player_id':'A1'})['observation_maps']['sources'][0]['estimated_person_seconds'] is None
    assert 'report.json' in excel_cell('x'*40000)


def test_geometry_diagnostics_visible_and_stale_verified_mask_rejected(project,tmp_path):
    project.put('calibration',{'id':'stale','source_id':'source','source_sha256':'different-video',
        'verified':True,'geometry':'court','matrix':[[1,0,0],[0,1,0],[0,0,1]],'visible_floor':[[0,0],[18,0],[18,9],[0,9]]})
    assert validated_court_profiles(project)==[]
    project.put('geometry_warning',{'id':'warn','source_id':'source','time':.5,'state':'shifted','reason':'Проверка сдвига'})
    project.put('camera_session',{'id':'session','source_id':'source','start':.5,'end':1,'geometry':'video','reason':'Новый ракурс'})
    out=export_all(project,tmp_path/'diagnostics')
    text=(out/'report.html').read_text(encoding="utf-8")
    assert 'Проверка сдвига' in text and 'viewer.html#source:0.500000' in text and 'Новый ракурс' in text
    assert '<polygon' not in (out/'court.svg').read_text(encoding="utf-8")
    from openpyxl import load_workbook
    book=load_workbook(out/'statistics.xlsx')
    assert {'Геометрия','Сессии камеры','Временная модель'}<=set(book.sheetnames)
